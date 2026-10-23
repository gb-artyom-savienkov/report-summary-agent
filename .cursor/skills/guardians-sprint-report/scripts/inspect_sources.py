#!/usr/bin/env python3
"""Preflight one sprint's local sources and write a compact reusable manifest."""

from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path

from docx import Document
from pdfminer.pdfparser import PDFSyntaxError

from extract_jira_snapshot import ROOT, find_snapshot
from render_report import CHARTS, inspect_pdf


def inspect_template(path: Path) -> list[str]:
    if not path.is_file():
        return [f"Template missing: {path}"]
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    required = ("{sprint_n}", "{sprint_start}", "{sprint_end}", "{sprint_summary}",
                "{velocity}", "{sa}", *CHARTS)
    errors = [f"Template placeholder missing: {item}" for item in required if item not in text]
    if len(doc.tables) != 2 or any(len(table.columns) != 2 or len(table.rows) < 2 for table in doc.tables):
        errors.append("Template tables changed; expected two two-column tables with a sample body row")
    return errors


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sprint", type=int, required=True)
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--out", type=Path)
    args = parser.parse_args()
    if args.sprint < 1:
        parser.error("sprint must be a positive integer")
    root = args.root.resolve()
    source_dir = root / f"source/sprint_{args.sprint}"
    snapshot, snapshot_warnings, snapshot_errors = find_snapshot(root / "source", args.sprint)
    errors = snapshot_errors + inspect_template(root / "template/Guardians - Sprint XXX Summary Report - template.docx")
    pdf = source_dir / "R&D Execution and Quality.pdf"
    try:
        pdf_info = inspect_pdf(pdf, args.sprint)
    except (ValueError, OSError, PDFSyntaxError) as exc:
        errors.append(str(exc))
        pdf_info = {"capture_date": None, "charts": {}, "warnings": []}
    gaps = []
    if snapshot is None:
        gaps.append("Matching saved Jira sprint report HTML unavailable; historical classification needs Jira change history or a matching export")
    if not pdf.is_file():
        gaps.append("Execution and quality PDF unavailable; KPIs and charts need a current export")
    elif len(pdf_info["charts"]) < len(CHARTS):
        gaps.append("One or more required PDF charts are absent; provide a complete export or use labeled placeholders")
    if snapshot and snapshot.get("snapshot_completion_display") and pdf_info["capture_date"]:
        try:
            closed = datetime.strptime(snapshot["snapshot_completion_display"], "%d/%b/%y %I:%M %p").date()
        except ValueError:
            snapshot_warnings.append("Jira snapshot completion date format changed; verify PDF freshness manually")
        else:
            if pdf_info["capture_date"] < closed:
                gaps.append(f"PDF predates Jira sprint completion ({closed}); provide a newer export")
    snapshot_summary = None
    if snapshot is not None:
        snapshot_summary = {key: snapshot[key] for key in (
            "sprint_number", "sprint_name", "sprint_id", "snapshot_start_display",
            "snapshot_completion_display", "section_counts")}
        snapshot_summary["snapshot_path"] = str(Path(snapshot["snapshot_path"]).relative_to(root))
        snapshot_summary["section_keys"] = {
            name: [item["key"] for item in rows] for name, rows in snapshot["sections"].items()}
    output = {
        "sprint_number": args.sprint,
        "status": "parser_change_needed" if errors else "source_gap" if gaps else "ready_for_review",
        "snapshot": snapshot_summary,
        "pdf": {
            "path": str(pdf.relative_to(root)),
            "capture_date": str(pdf_info["capture_date"]) if pdf_info["capture_date"] else None,
            "charts_found": sorted(pdf_info["charts"]),
            "chart_labels_required": {kind: label for label, kind in CHARTS.values()},
        },
        "errors": errors,
        "gaps": gaps,
        "warnings": snapshot_warnings + pdf_info["warnings"],
    }
    destination = args.out or root / f"dist/sprint_{args.sprint}/source-manifest.json"
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    draft = destination.parent / "report-data.draft.json"
    if snapshot is not None and not draft.exists() and not (destination.parent / "report-data.json").exists():
        row = lambda item: {"key": item["key"], "summary": item["summary"], "comment": "",
                            "snapshot_status": item["snapshot_status"],
                            "story_points_display": item["story_points_display"]}
        draft_data = {
            "sprint_number": args.sprint,
            "sprint_start": "",
            "sprint_end": "",
            "velocity": None,
            "sprint_accuracy": None,
            "classification_basis": "saved_jira_html",
            "classification_reference": snapshot_summary["snapshot_path"],
            "summary_paragraphs": [],
            "delivered": [row(item) for item in snapshot["sections"]["completed"]],
            "not_completed": [row(item) for item in snapshot["sections"]["not_completed"] + snapshot["sections"]["removed"]],
            "completed_outside": [{"key": item["key"], "summary": item["summary"]}
                                  for item in snapshot["sections"]["completed_outside"]],
            "material_discrepancies": [],
        }
        draft.write_text(json.dumps(draft_data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    summary = {
        "status": output["status"],
        "manifest": str(destination),
        "draft": str(draft) if draft.exists() else None,
        "sprint": args.sprint,
        "section_counts": snapshot["section_counts"] if snapshot else None,
        "charts_found": output["pdf"]["charts_found"],
        "errors": errors,
        "gaps": gaps,
        "warnings": output["warnings"],
    }
    print(json.dumps(summary, ensure_ascii=False))
    return 2 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
