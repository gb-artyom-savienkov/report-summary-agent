#!/usr/bin/env python3
"""Fill the branded Guardians DOCX from reviewed, sprint-specific JSON data."""

from __future__ import annotations

import argparse
import json
import logging
import re
import shutil
import subprocess
import sys
import tempfile
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

import pdfplumber
from pdfminer.pdfparser import PDFSyntaxError
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import _Row
from docx.shared import Inches, Pt
from PIL import Image

from extract_jira_snapshot import parse_snapshot


logging.getLogger("pdfminer").setLevel(logging.ERROR)

ROOT = Path(__file__).resolve().parents[4]
CHARTS = {
    "{velocity_chart}": ("Sprint Velocity (Current Quater)", "velocity"),
    "{escalated_cs_chart}": ("Customer Support Backlog", "customer_support"),
    "{bugs_backlog_chart}": ("Bug Backlog", "bug_backlog"),
}


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def source_path(reference: str, root: Path) -> Path:
    path = Path(reference)
    return path if path.is_absolute() else root / path


def validate(data: dict, root: Path) -> dict | None:
    number = data.get("sprint_number")
    require(type(number) is int and number > 0, "sprint_number must be a positive integer")
    for field in ("sprint_start", "sprint_end"):
        require(isinstance(data.get(field), str) and data[field].strip(), f"{field} is required")
    paragraphs = data.get("summary_paragraphs")
    require(isinstance(paragraphs, list) and len(paragraphs) in (4, 5), "summary_paragraphs must have 4 or 5 paragraphs")
    require(all(isinstance(p, str) and p.strip() for p in paragraphs), "summary paragraphs cannot be empty")
    require(all(not re.search(r"\[[^]]*\]", p) for p in paragraphs), "summary contains bracketed drafting placeholders")
    for field in ("velocity", "sprint_accuracy"):
        require(data.get(field) is None or isinstance(data[field], str), f"{field} must be a string or null")
    if data.get("velocity") and data.get("sprint_accuracy"):
        prefix = f"Sprint {number} delivered a **velocity of {data['velocity']}** with a **sprint accuracy of {data['sprint_accuracy']}**."
        require(paragraphs[0].startswith(prefix), "summary must begin with the exact reported KPIs")
    basis = data.get("classification_basis")
    require(basis in ("saved_jira_html", "jira_retrospective", "jira_history"), "unsupported classification_basis")
    reference = data.get("classification_reference")
    require(isinstance(reference, str) and reference.strip(), "classification_reference is required")
    delivered = data.get("delivered")
    incomplete = data.get("not_completed")
    require(isinstance(delivered, list) and isinstance(incomplete, list), "delivered and not_completed must be arrays")
    seen: set[str] = set()
    for section, rows in (("delivered", delivered), ("not_completed", incomplete)):
        for row in rows:
            require(isinstance(row, dict), f"{section} rows must be objects")
            key = row.get("key", "")
            require(bool(re.fullmatch(r"GNG-\d+", key)), f"invalid issue key: {key!r}")
            require(key not in seen, f"duplicate issue key: {key}")
            seen.add(key)
            require(isinstance(row.get("summary"), str) and row["summary"].strip(), f"missing summary for {key}")
            require(isinstance(row.get("comment", ""), str), f"comment must be a string for {key}")
    if basis == "saved_jira_html":
        path = source_path(reference, root)
        require(path.is_file(), f"saved Jira HTML not found: {path}")
        snapshot = parse_snapshot(path)
        require(snapshot["sprint_number"] == number, f"saved HTML is Sprint {snapshot['sprint_number']}, not Sprint {number}")
        expected_delivered = {r["key"] for r in snapshot["sections"]["completed"]}
        expected_incomplete = {r["key"] for r in snapshot["sections"]["not_completed"] + snapshot["sections"]["removed"]}
        require({r["key"] for r in delivered} == expected_delivered, "Delivered keys do not match saved Jira Completed Issues")
        require({r["key"] for r in incomplete} == expected_incomplete, "Not completed keys must match saved Jira Not Completed and Removed sections")
        for section in snapshot["sections"].values():
            for row in section:
                if row["key"] in seen:
                    require(valid_issue_url(row.get("url"), row["key"]), f"missing or invalid HTML ticket URL for {row['key']}")
        return snapshot
    return None


def valid_issue_url(url: str | None, key: str) -> bool:
    if not isinstance(url, str):
        return False
    parsed = urlparse(url)
    return parsed.scheme == "https" and parsed.netloc == "glassboxdigital.atlassian.net" and parsed.path == f"/browse/{key}"


def issue_urls(data: dict, snapshot: dict | None) -> dict[str, str]:
    if data["classification_basis"] == "saved_jira_html":
        require(snapshot is not None, "saved Jira snapshot was not loaded")
        return {row["key"]: row["url"] for section in snapshot["sections"].values() for row in section}
    urls = {row["key"]: row.get("url") for row in data["delivered"] + data["not_completed"]}
    for key, url in urls.items():
        require(valid_issue_url(url, key), f"missing or invalid Jira ticket URL for {key}")
    return urls


def replace_in_runs(paragraph, mapping: dict[str, str]) -> None:
    for old, new in mapping.items():
        while old in paragraph.text:
            start = paragraph.text.index(old)
            end = start + len(old)
            offset = 0
            first = last = None
            for index, run in enumerate(paragraph.runs):
                next_offset = offset + len(run.text)
                if first is None and start < next_offset:
                    first = (index, start - offset)
                if first is not None and end <= next_offset:
                    last = (index, end - offset)
                    break
                offset = next_offset
            require(first is not None and last is not None, f"cannot replace {old} in paragraph")
            first_index, first_offset = first
            last_index, last_offset = last
            runs = paragraph.runs
            if first_index == last_index:
                run = runs[first_index]
                run.text = run.text[:first_offset] + new + run.text[last_offset:]
            else:
                suffix = runs[last_index].text[last_offset:]
                runs[first_index].text = runs[first_index].text[:first_offset] + new
                for index in range(first_index + 1, last_index):
                    runs[index].text = ""
                runs[last_index].text = suffix


def fill_summary(paragraph, paragraphs: list[str]) -> None:
    sample_rpr = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    current = paragraph
    for index, content in enumerate(paragraphs):
        if index:
            new_xml = deepcopy(paragraph._p)
            current._p.addnext(new_xml)
            current = Paragraph(new_xml, paragraph._parent)
        current.clear()
        current.paragraph_format.space_after = Pt(8)
        for part in re.split(r"(\*\*[^*]+\*\*)", content):
            if not part:
                continue
            bold = part.startswith("**") and part.endswith("**")
            run = current.add_run(part[2:-2] if bold else part)
            if sample_rpr is not None:
                run._r.insert(0, deepcopy(sample_rpr))
            if bold:
                run.bold = True


def set_cell(cell, content: str) -> None:
    first = cell.paragraphs[0]
    sample_rpr = deepcopy(first.runs[0]._r.rPr) if first.runs and first.runs[0]._r.rPr is not None else None
    first.clear()
    run = first.add_run(content)
    if sample_rpr is not None:
        run._r.insert(0, sample_rpr)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def set_table_font(rpr) -> None:
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), "Calibri")
    for size_name in ("w:sz", "w:szCs"):
        size = rpr.find(qn(size_name))
        if size is None:
            size = OxmlElement(size_name)
            rpr.append(size)
        size.set(qn("w:val"), "22")


def set_issue_cell(cell, key: str, summary: str, url: str | None) -> None:
    require(len(cell.paragraphs) >= 2, "template issue cell must have key and title paragraphs")
    key_paragraph, title_paragraph = cell.paragraphs[:2]
    key_run = next((run for run in key_paragraph.runs if run.text.strip()), None)
    title_run = next((run for run in title_paragraph.runs if run.text.strip()), None)
    key_rpr = deepcopy(key_run._r.rPr) if key_run is not None and key_run._r.rPr is not None else None
    title_rpr = deepcopy(title_run._r.rPr) if title_run is not None and title_run._r.rPr is not None else None
    key_paragraph.clear()
    title_paragraph.clear()
    if key:
        require(url is not None and valid_issue_url(url, key), f"invalid issue URL for {key}")
        relationship_id = key_paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), relationship_id)
        run = OxmlElement("w:r")
        if key_rpr is not None:
            run.append(key_rpr)
        run_properties = run.find(qn("w:rPr"))
        if run_properties is None:
            run_properties = OxmlElement("w:rPr")
            run.insert(0, run_properties)
        set_table_font(run_properties)
        color = run_properties.find(qn("w:color"))
        if color is None:
            color = OxmlElement("w:color")
            run_properties.append(color)
        color.set(qn("w:val"), "1868DB")
        underline = run_properties.find(qn("w:u"))
        if underline is None:
            underline = OxmlElement("w:u")
            run_properties.append(underline)
        underline.set(qn("w:val"), "single")
        issue_text = OxmlElement("w:t")
        issue_text.text = key
        run.append(issue_text)
        hyperlink.append(run)
        key_paragraph._p.append(hyperlink)
    title = title_paragraph.add_run(summary)
    if title_rpr is not None:
        title._r.insert(0, title_rpr)
    title.font.name = "Calibri"
    title.font.size = Pt(11)
    for extra in cell.paragraphs[2:]:
        extra._element.getparent().remove(extra._element)


def fill_table(table, rows: list[dict], empty_text: str, urls: dict[str, str]) -> None:
    require(len(table.columns) == 2 and len(table.rows) >= 2, "template table structure changed")
    sample_xml = table.rows[1]._tr
    display = rows if rows else [{"key": "", "summary": empty_text, "comment": ""}]
    insert_after = sample_xml
    for item in display:
        clone = deepcopy(sample_xml)
        insert_after.addnext(clone)
        insert_after = clone
        row = _Row(clone, table)
        set_issue_cell(row.cells[0], item["key"], item["summary"], urls.get(item["key"]))
        set_cell(row.cells[1], item.get("comment", ""))
    sample_xml.getparent().remove(sample_xml)
    for index, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        if index == 0 and tr_pr.find(qn("w:tblHeader")) is None:
            tr_pr.append(OxmlElement("w:tblHeader"))
        if index > 0:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    ppr = paragraph._p.get_or_add_pPr()
                    paragraph_rpr = ppr.find(qn("w:rPr"))
                    if paragraph_rpr is None:
                        paragraph_rpr = OxmlElement("w:rPr")
                        ppr.append(paragraph_rpr)
                    set_table_font(paragraph_rpr)
                    for run_rpr in paragraph._p.xpath(".//w:r/w:rPr"):
                        set_table_font(run_rpr)


def chart_box(kind: str, page, words: list[dict]) -> tuple[float, float, float, float] | None:
    height, width = page.height, page.width
    if kind == "velocity":
        anchor = next((w for w in words if w["text"] == "Quater)" and w["top"] > height * .5), None)
        if not anchor or "Sprint Velocity (Current Quater)" not in (page.extract_text() or ""):
            return None
        return (width * .505, anchor["top"] - 37, width * .986, anchor["top"] + 315)
    if kind == "customer_support":
        anchors = [w for w in words if w["text"] == "Backlog" and w["top"] < height * .25]
        anchor = next((w for w in anchors if width * .4 < w["x0"] < width * .55), None)
        if not anchor or "Customer Support Backlog" not in (page.extract_text() or ""):
            return None
        return (width * .33, anchor["top"] - 20, width * .648, anchor["top"] + 305)
    if kind == "bug_backlog":
        anchors = [w for w in words if w["text"] == "BugBacklog" and height * .45 < w["top"] < height * .8]
        anchor = next((w for w in anchors if w["x0"] < width * .15), None)
        if not anchor or "Bug Backlog" not in (page.extract_text() or ""):
            return None
        return (width * .018, anchor["top"] - 5, width * .49, anchor["top"] + 315)
    return None


def inspect_pdf(pdf: Path, sprint_number: int) -> dict:
    """Read the PDF once and classify missing panels versus changed layouts."""
    result = {"capture_date": None, "charts": {}, "warnings": []}
    if not pdf.is_file():
        result["warnings"].append(f"PDF missing: {pdf}")
        return result
    with pdfplumber.open(pdf) as document:
        raw_date = (document.metadata or {}).get("CreationDate", "")
        match = re.match(r"D:(\d{4})(\d{2})(\d{2})", raw_date)
        if match:
            result["capture_date"] = datetime(*map(int, match.groups())).date()
        page_texts = [page.extract_text() or "" for page in document.pages]
        require(any(page_texts), f"PDF has no extractable text: {pdf}; provide a text-based export or update the parser")
        for _, (label, kind) in CHARTS.items():
            matching_pages = [(index, document.pages[index]) for index, content in enumerate(page_texts) if label in content]
            if not matching_pages:
                hint = {"velocity": "Sprint Velocity", "customer_support": "Customer Support", "bug_backlog": "Bug Backlog"}[kind]
                require(not any(hint in content for content in page_texts),
                        f"PDF heading for {kind} changed; update CHARTS/chart_box() after inspecting the export")
                result["warnings"].append(f"Missing PDF panel: {label}")
                continue
            candidates = []
            for index, page in matching_pages:
                box = chart_box(kind, page, page.extract_words())
                if box:
                    candidates.append((index + 1, page.width, page.height, box, page))
            require(candidates, f"PDF layout changed for {label}; update chart_box() before rendering")
            require(len(candidates) == 1, f"PDF has {len(candidates)} {label} panels; select the correct panel before rendering")
            page_number, width, height, box, page = candidates[0]
            if kind == "velocity":
                cropped_text = page.crop(box).extract_text() or ""
                require(re.search(rf"NG\s*Sprint\s*{sprint_number}\b", cropped_text, re.I) is not None,
                        f"velocity panel has no Sprint {sprint_number} row; provide the correct PDF")
            if kind == "bug_backlog" and "Quality Score" in page_texts[page_number - 1]:
                result["warnings"].append("Bug Backlog is under Quality Score in the PDF; expected Execution Score")
            result["charts"][kind] = (page_number, width, height, box)
    return result


def extract_charts(pdf: Path, charts: dict, temp: Path) -> dict[str, Path]:
    """Rasterize each source page once, then crop its requested panels."""
    if not charts:
        return {}
    poppler = shutil.which("pdftoppm")
    require(poppler is not None, "pdftoppm is required to extract charts; install Poppler")
    images: dict[int, Path] = {}
    for page_number in sorted({spec[0] for spec in charts.values()}):
        prefix = temp / f"page-{page_number}"
        subprocess.run([poppler, "-f", str(page_number), "-l", str(page_number), "-r", "180", "-png", "-singlefile", str(pdf), str(prefix)],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        images[page_number] = prefix.with_suffix(".png")
    output = {}
    for kind, (page_number, width, height, box) in charts.items():
        with Image.open(images[page_number]) as image:
            scaled = tuple(int(value * (image.width / width if index % 2 == 0 else image.height / height))
                           for index, value in enumerate(box))
            output[kind] = temp / f"{kind}.png"
            image.crop(scaled).save(output[kind])
    return output


def insert_chart(paragraph, chart: Path | None, label: str) -> None:
    paragraph.clear()
    if chart is None:
        paragraph.add_run(f"Chart unavailable: {label}")
        return
    run = paragraph.add_run()
    run.add_picture(str(chart), width=Inches(6.4))


def sprint_end_date(display: str):
    for pattern in ("%b %d, %Y", "%B %d, %Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(display, pattern).date()
        except ValueError:
            pass
    return None


def normalize_section_break(doc, heading: str) -> None:
    paragraph = next((p for p in doc.paragraphs if p.text == heading), None)
    if paragraph is None:
        raise ValueError(f"section heading missing: {heading}")
    previous = paragraph._p.getprevious()
    removed_break = False
    while previous is not None and previous.tag == qn("w:p") and not "".join(previous.itertext()).strip():
        removed_break = removed_break or 'w:type="page"' in previous.xml
        old = previous
        previous = previous.getprevious()
        old.getparent().remove(old)
    if removed_break:
        paragraph.paragraph_format.page_break_before = True


def validate_document(doc, data: dict, urls: dict[str, str]) -> None:
    all_text = "\n".join(p.text for p in doc.paragraphs)
    all_text += "\n" + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    require(not re.search(r"\{(?:sprint_|velocity|sa|escalated_cs_chart|bugs_backlog_chart)", all_text),
            "unreplaced template variable remains")
    require("Source PDF places Bug Backlog under Quality Score" not in all_text,
            "source layout note leaked into report body")
    require(len(doc.tables) == 2, "report table count changed")
    for table, items in zip(doc.tables, (data["delivered"], data["not_completed"])):
        require(len(table.rows) == len(items) + 1 if items else len(table.rows) == 2,
                "report table row count incorrect")
        for row, item in zip(table.rows[1:], items):
            key = item["key"]
            key_paragraph, title_paragraph = row.cells[0].paragraphs[:2]
            require(key_paragraph.text == key and title_paragraph.text == item["summary"],
                    f"issue cell text changed for {key}")
            hyperlinks = key_paragraph._p.xpath("./w:hyperlink")
            require(len(hyperlinks) == 1, f"missing Jira hyperlink for {key}")
            relationship_id = hyperlinks[0].get(qn("r:id"))
            require(key_paragraph.part.rels[relationship_id].target_ref == urls[key],
                    f"Jira hyperlink target changed for {key}")
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    require(paragraph.paragraph_format.line_spacing == 1.0,
                            f"table line spacing changed for {key}")
                    rprs = [paragraph._p.pPr.find(qn("w:rPr")), *paragraph._p.xpath(".//w:r/w:rPr")]
                    for rpr in rprs:
                        require(rpr is not None, f"table font missing for {key}")
                        fonts = rpr.find(qn("w:rFonts"))
                        require(fonts is not None and all(fonts.get(qn(f"w:{name}")) == "Calibri"
                                                           for name in ("ascii", "hAnsi", "eastAsia", "cs")),
                                f"table font is not Calibri for {key}")
                        require(all(rpr.find(qn(name)) is not None and rpr.find(qn(name)).get(qn("w:val")) == "22"
                                    for name in ("w:sz", "w:szCs")),
                                f"table font size is not 11 pt for {key}")


def render(data: dict, root: Path, output: Path, snapshot: dict | None) -> list[str]:
    template = root / "template/Guardians - Sprint XXX Summary Report - template.docx"
    pdf = root / f"source/sprint_{data['sprint_number']}/R&D Execution and Quality.pdf"
    require(template.is_file(), f"template not found: {template}")
    doc = Document(template)
    mapping = {
        "{sprint_n}": str(data["sprint_number"]),
        "{sprint_start}": data["sprint_start"],
        "{sprint_end}": data["sprint_end"],
        "{velocity}": data.get("velocity") or "",
        "~{sa}%": data.get("sprint_accuracy") or "",
        "{sa}": data.get("sprint_accuracy") or "",
    }
    summary = next((p for p in doc.paragraphs if "{sprint_summary}" in p.text), None)
    require(summary is not None, "summary placeholder missing from template")
    fill_summary(summary, data["summary_paragraphs"])
    normalize_section_break(doc, "Sprint Delivery")
    normalize_section_break(doc, "Sprint Reports")
    for paragraph in doc.paragraphs:
        replace_in_runs(paragraph, mapping)
    if not data.get("velocity") or not data.get("sprint_accuracy"):
        metric = next((p for p in doc.paragraphs if "Velocity:" in p.text or "Sprint accuracy:" in p.text), None)
        require(metric is not None, "metric paragraph missing from template")
        metric.clear()
        if data.get("velocity"):
            metric.add_run(f"Velocity: {data['velocity']}")
        if data.get("sprint_accuracy"):
            if data.get("velocity"):
                metric.add_run("\n")
            metric.add_run(f"Sprint accuracy: {data['sprint_accuracy']}")
        if not data.get("velocity") and not data.get("sprint_accuracy"):
            note = metric.add_run(f"Sprint {data['sprint_number']} velocity and sprint accuracy were not reported in the supplied PDF.")
            note.italic = True
    urls = issue_urls(data, snapshot)
    fill_table(doc.tables[0], data["delivered"], "No completed issues recorded in the source", urls)
    fill_table(doc.tables[1], data["not_completed"], "No issues listed as not completed or removed in the source", urls)
    for heading_text in ("Escalated Customer ticket backlog", "Bugs backlog"):
        chart_heading = next((p for p in doc.paragraphs if p.text.strip() == heading_text), None)
        require(chart_heading is not None, f"{heading_text} heading missing from template")
        chart_heading.paragraph_format.keep_with_next = True
        if heading_text == "Escalated Customer ticket backlog" and data.get("page_break_before_customer_support_chart"):
            chart_heading.paragraph_format.page_break_before = True
    inspection = inspect_pdf(pdf, data["sprint_number"])
    warnings = inspection["warnings"]
    capture_date = inspection["capture_date"]
    end_date = sprint_end_date(data["sprint_end"])
    stale_pdf = capture_date is not None and end_date is not None and capture_date < end_date
    if stale_pdf:
        require(not data.get("velocity") and not data.get("sprint_accuracy"),
                "a PDF created before sprint close cannot supply final velocity or sprint accuracy")
        warnings.append(f"Source PDF was created {capture_date}, before Sprint {data['sprint_number']} ended {end_date}; all charts marked unavailable")
    elif not pdf.is_file():
        require(not data.get("velocity") and not data.get("sprint_accuracy"),
                "missing source PDF cannot supply velocity or sprint accuracy")
    elif capture_date is None:
        warnings.append("PDF creation date unavailable; verify that it was exported after sprint close")
    with tempfile.TemporaryDirectory(prefix="guardians-charts-") as tmp:
        temp = Path(tmp)
        chart_images = {} if stale_pdf else extract_charts(pdf, inspection["charts"], temp)
        for placeholder, (label, kind) in CHARTS.items():
            p = next((p for p in doc.paragraphs if placeholder in p.text), None)
            require(p is not None, f"chart placeholder {placeholder} missing from template")
            chart = chart_images.get(kind)
            insert_chart(p, chart, label)
            if chart is None:
                warnings.append(f"{label}: chart unavailable")
        while doc.paragraphs and not doc.paragraphs[-1].text.strip() and not doc.paragraphs[-1]._p.xpath('.//w:drawing'):
            trailing = doc.paragraphs[-1]._p
            trailing.getparent().remove(trailing)
        output.parent.mkdir(parents=True, exist_ok=True)
        staging = output.with_name(output.stem + ".staging.docx")
        try:
            doc.save(staging)
            reopened = Document(staging)
            validate_document(reopened, data, urls)
            staging.replace(output)
        finally:
            staging.unlink(missing_ok=True)
    return warnings


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("data", type=Path)
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--out", type=Path)
    args = parser.parse_args()
    try:
        data = json.loads(args.data.read_text(encoding="utf-8"))
        root = args.root.resolve()
        snapshot = validate(data, root)
        output = args.out or root / f"dist/sprint_{data['sprint_number']}/Guardians - Sprint {data['sprint_number']} Summary Report.docx"
        warnings = render(data, root, output, snapshot)
    except (ValueError, OSError, PDFSyntaxError, json.JSONDecodeError, subprocess.CalledProcessError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 2
    print(output.resolve())
    for warning in warnings:
        print(f"Warning: {warning}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
