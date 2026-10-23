"""Regression checks for source identity, changed layouts, and portable rendering."""

import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / ".cursor/skills/guardians-sprint-report/scripts"))

from extract_jira_snapshot import parse_snapshot  # noqa: E402
from render_report import inspect_pdf, replace_in_runs, source_path  # noqa: E402


class PipelineTests(unittest.TestCase):
    def test_placeholders_can_span_template_runs(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("Sprint {")
        paragraph.add_run("sprint_n")
        paragraph.add_run("} Summary")
        replace_in_runs(paragraph, {"{sprint_n}": "124"})
        self.assertEqual(paragraph.text, "Sprint 124 Summary")

    def test_relative_source_reference_uses_repository_root(self):
        self.assertEqual(source_path("source/sprint_124/report.html", ROOT),
                         ROOT / "source/sprint_124/report.html")

    def test_unknown_jira_section_requires_parser_update(self):
        markup = '''<select id="ghx-chart-picker"><option selected value="4">NG Sprint 124</option></select>
        <div class="ghx-sprint-report-table" data-sprint-id="4"><table aria-label="Renamed Issues">
        <tr><td><a href="https://glassboxdigital.atlassian.net/browse/GNG-1">GNG-1</a></td>
        <td>Work</td><td>Bug</td><td>High</td><td>Done</td><td>1</td></tr></table></div>'''
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "report.html"
            path.write_text(markup, encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "unknown Jira report section"):
                parse_snapshot(path)

    def test_wrong_sprint_pdf_is_rejected(self):
        pdf = ROOT / "source/sprint_123/R&D Execution and Quality.pdf"
        if not pdf.exists():
            self.skipTest("Sprint 123 PDF fixture is absent")
        with self.assertRaisesRegex(ValueError, "no Sprint 124 row"):
            inspect_pdf(pdf, 124)


if __name__ == "__main__":
    unittest.main()
